#!/usr/bin/env python3
"""
notice_reconcile.py
===================
1. Reads every Section-138 notice PDF in a folder.
2. Extracts: agreement no, cheque no/date/amount/bank, return reason and every
   party (name + address) -- including parties that have several addresses
   ("Also at ...").
3. Compares them with company.xlsx (wide format: Borrower, Co-Borrower-1 ...)
   and writes reconciliation_report.xlsx with an OK / CHECK verdict per notice.

USAGE
-----
    python notice_reconcile.py                       # PDFs + company.xlsx in this folder
    python notice_reconcile.py --pdf-folder D:\\notices --company D:\\company.xlsx

INSTALL
-------
    pip install pandas openpyxl pdfplumber rapidfuzz     (rapidfuzz optional)
"""

import argparse
import os
import re
import sys
import time
from datetime import date, datetime

import pandas as pd

try:
    import pdfplumber
except ImportError:  # pragma: no cover
    pdfplumber = None

# ----------------------------------------------------------------------------
# CONFIG
# ----------------------------------------------------------------------------
MATCH_T = 90      # name AND address >= this      -> MATCH
PARTIAL_T = 60    # name AND address >= this      -> PARTIAL MATCH, else MISMATCH
PAIR_MIN_NAME = 50  # a notice party is only paired with a company party if the
                    # names are at least this similar (else NOT IN COMPANY DATA)

# ----------------------------------------------------------------------------
# Fuzzy similarity (rapidfuzz if present, difflib otherwise)
# ----------------------------------------------------------------------------
try:
    from rapidfuzz import fuzz

    def _ratio(a, b):
        return fuzz.ratio(a, b)

    def _tsort(a, b):
        return fuzz.token_sort_ratio(a, b)

except ImportError:
    from difflib import SequenceMatcher

    def _ratio(a, b):
        return 100.0 * SequenceMatcher(None, a, b).ratio()

    def _tsort(a, b):
        return _ratio(" ".join(sorted(a.split())), " ".join(sorted(b.split())))


def isnull(v):
    if v is None:
        return True
    try:
        return bool(pd.isna(v))
    except (TypeError, ValueError):
        return False


def norm(text):
    """Upper-case, punctuation -> space, collapse whitespace."""
    if isnull(text):
        return ""
    t = re.sub(r"[^A-Z0-9]+", " ", str(text).upper())
    return t.strip()


def squash(text):
    return norm(text).replace(" ", "")


def similarity(a, b):
    """0-100. Ignores spaces (company data often glues words together) and word order."""
    a, b = norm(a), norm(b)
    if not a and not b:
        return 100.0
    if not a or not b:
        return 0.0
    return max(_ratio(a.replace(" ", ""), b.replace(" ", "")), _tsort(a, b))


def parse_date(v):
    if isnull(v) or str(v).strip() == "":
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v).strip()
    m = re.match(r"^(\d{1,2})[-/.](\d{1,2})[-/.](\d{2,4})$", s)
    if m:
        d, mo, y = (int(x) for x in m.groups())
        if y < 100:
            y += 2000
        try:
            return date(y, mo, d)
        except ValueError:
            return None
    try:
        return pd.to_datetime(s, dayfirst=True).date()
    except Exception:
        return None


def to_number(v):
    if isnull(v):
        return None
    s = re.sub(r"[^\d.]", "", str(v))
    try:
        return float(s) if s else None
    except ValueError:
        return None


def cheque_key(v):
    """'000504' == 504 == 504.0 == '504'"""
    if isnull(v):
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return re.sub(r"[^A-Z0-9]", "", str(v).upper()).lstrip("0")


# ----------------------------------------------------------------------------
# PART 1 - PDF EXTRACTION
# ----------------------------------------------------------------------------
DATE_RX = r"\d{1,2}[-/.]\d{1,2}[-/.]\d{2,4}"

# "Also at", "Also residing at", "Also Address:", "Alternate address" ...
ALSO_RE = re.compile(
    r"\b(?:also\s+(?:at|add(?:ress)?|residing\s+at|resides\s+at|r/o)"
    r"|alternate\s+address|another\s+address)\b\s*[:\-\u2013\u2014.]*\s*",
    re.IGNORECASE,
)
NUM_RE = re.compile(r"^(\d+)\s*[.)]\s*(.*)$")
NOISE_RES = [
    re.compile(r"^[A-Z]\s*[A-Z]\s*(?:\d\s*){6,}[A-Z]\s*[A-Z]$", re.I),  # speed-post barcode no.
    re.compile(r"^DATED\s*:", re.I),
    re.compile(r"^(?:by\s+)?(?:speed\s+post|registered\s+post|regd\.?\s*a\.?d\.?)\b", re.I),
]

TO_RE = re.compile(
    r"^[ \t]*To\b[ \t]*[,:\-]?[ \t]*(.*?)"
    r"(?=^[ \t]*(?:Dear\s+Sir|Sub(?:ject)?\s*[:\-]|Under\s+the\s+instructions))",
    re.IGNORECASE | re.DOTALL | re.MULTILINE,
)
AGREEMENT_RES = [
    re.compile(
        r"(?:loan\s+account\s+no|loan\s+a/?c\s+no|loan\s+agreement\s+no|agreement\s+no)"
        r"\s*\.?\s*:?\s*-?\s*([A-Z0-9]{8,})",
        re.I,
    ),
    re.compile(r"\b([A-Z]{2}\d{2}[A-Z]{3}\d{8,})\b"),
]


def is_noise(line):
    return any(rx.search(line) for rx in NOISE_RES)


LETTERHEAD_RE = re.compile(
    r"[ \t]*Registered\s+Office\s*:.*?Branches\s*:[^\n]*\n?",
    re.IGNORECASE | re.DOTALL,
)
FIRM_NAME_RE = re.compile(
    r"^[ \t]*Z\.?\s*SANWARWALA\s*&?\s*COMPANY[ \t]*\n[ \t]*Advocates[ \t]*\n?",
    re.IGNORECASE | re.MULTILINE,
)


def strip_letterhead(text):
    """
    The firm's letterhead (name, 'Advocates', Registered/Corporate Office,
    Branches) reprints at the top of every page. When a sentence (e.g. the
    cheque details) happens to start a new page, the letterhead lands in the
    middle of it and breaks the regexes. Drop every occurrence - it carries
    no case data.
    """
    text = FIRM_NAME_RE.sub("", text)
    text = LETTERHEAD_RE.sub("", text)
    return text


def read_pdf_text(path):
    if pdfplumber is None:
        sys.exit("pdfplumber is missing:  pip install pdfplumber")
    chunks = []
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            t = page.extract_text()
            if t:
                chunks.append(t)
    text = "\n".join(chunks).replace("\r\n", "\n").replace("\r", "\n")
    return strip_letterhead(text)


def parse_recipients(block):
    """
    Returns a list of parties: {"num", "name", "addresses": [str, ...]}.
    - Party numbers must run 1, 2, 3 ... in order, so an address line such as
      "5. Main Road" can never be mistaken for a new party.
    - "Also at ..." (anywhere in a line) starts an additional address for the
      *same* party.
    """
    parties, cur, expected = [], None, 1
    for raw in block.split("\n"):
        line = raw.strip()
        if not line or is_noise(line):
            continue
        m = NUM_RE.match(line)
        if m and int(m.group(1)) == expected:
            cur = {"num": expected, "name": m.group(2).strip(), "addr_lines": [[]]}
            parties.append(cur)
            expected += 1
            continue
        if cur is None:
            continue
        if not cur["name"]:            # number sat alone on its line
            cur["name"] = line
            continue
        if re.match(r"^\d+\s*[.)]?$", line):   # stray list number
            continue
        first, *rest = ALSO_RE.split(line)
        if first.strip():
            cur["addr_lines"][-1].append(first.strip())
        for part in rest:
            cur["addr_lines"].append([part.strip()] if part.strip() else [])

    out = []
    for p in parties:
        addrs = [", ".join(a) for a in p["addr_lines"]]
        addrs = [a for i, a in enumerate(addrs) if a or i == 0]  # drop empty extras
        if p["name"] and len(p["name"]) >= 2:
            out.append({"num": p["num"], "name": p["name"], "addresses": addrs})
    return out


def parse_cheque(text):
    info = {"Cheque No": "", "Cheque Date": "", "Cheque Amount": "", "Drawn On": ""}
    para = re.search(
        r"cheque\s+(?:bearing\s+)?no\.?\s*[:\-]?\s*([\w\-/]+)\s*,?\s*dated\s*[-:\s]*(" + DATE_RX + r")"
        r"\s+for\s+a\s+sum\s+of\s+(?:Rs\.?|INR|\u20b9)\s*([\d,]+(?:\.\d+)?)(?:/-)?"
        r"\s+drawn\s+on\s+(.*?)(?=\s+in\s+favou?r|\s+vide|\.(?:\s|$)|\n|$)",
        text,
        re.I | re.S,
    )
    if para:
        info["Cheque No"], info["Cheque Date"], amt, info["Drawn On"] = (
            g.strip() for g in para.groups()
        )
        info["Cheque Amount"] = amt.replace(",", "")
        return info

    # fallback: pipe / tabular layout
    t = re.sub(r"\n\s*\|\s*", " | ", text)
    m = re.search(r"(" + DATE_RX + r")\s*\|\s*(\w+)\s*\|\s*([^|]+?)\s*(?:\||\n)", t)
    if m:
        info["Cheque Date"], info["Cheque No"], info["Drawn On"] = (g.strip() for g in m.groups())
    m = re.search(r"(?:TOTAL|AMOUNT)\s*[:\-]?\s*([\d,]+)", text, re.I)
    if m:
        info["Cheque Amount"] = m.group(1).replace(",", "")
    return info


def parse_notice_text(text, filename=""):
    """Pure-text parser (so it can be unit-tested without a PDF)."""
    agreement = ""
    for rx in AGREEMENT_RES:
        m = rx.search(text) or rx.search(filename)
        if m:
            agreement = m.group(1).upper()
            break

    cheque = parse_cheque(text)
    m = re.search(r"DATED\s*:?\s*(" + DATE_RX + ")", text, re.I)
    notice_date = m.group(1) if m else ""
    m = re.search(r"memo\s+dated\s*(" + DATE_RX + ")", text, re.I)
    memo_date = m.group(1) if m else ""
    m = re.search(r"endorsement\s*[\u201c\"'`]*\s*([^\u201d\"'`]+?)\s*[\u201d\"'`]", text, re.I | re.S)
    reason = re.sub(r"\s+", " ", m.group(1)).strip().rstrip(".") if m else ""

    parties = []
    to = TO_RE.search(text)
    if to:
        parties = parse_recipients(to.group(1))

    records = []
    for p in parties:
        for seq, addr in enumerate(p["addresses"], start=1):
            records.append(
                {
                    "Agreement No": agreement,
                    "Party Number": p["num"],
                    "Address Seq": seq,
                    "Name": p["name"],
                    "Address": addr,
                    "Cheque No": cheque["Cheque No"],
                    "Cheque Date": cheque["Cheque Date"],
                    "Cheque Amount": cheque["Cheque Amount"],
                    "Drawn On": cheque["Drawn On"],
                    "Return Reason": reason,
                    "Memo Date": memo_date,
                    "Notice Date": notice_date,
                    "Source File": os.path.basename(filename),
                }
            )
    return records


def extract_data_from_pdf(pdf_path):
    return parse_notice_text(read_pdf_text(pdf_path), pdf_path)


# ----------------------------------------------------------------------------
# PART 2 - COMPANY DATA (wide -> list of parties per agreement)
# ----------------------------------------------------------------------------
def load_company(path):
    df = pd.read_excel(path)
    cols = {re.sub(r"[^A-Z0-9]", "", str(c).upper()): c for c in df.columns}

    def pick(*names):
        for n in names:
            if n in cols:
                return cols[n]
        return None

    c_agr = pick("AGREEMENTNO", "AGREEMENTNUMBER", "LOANACCOUNTNO", "LOANACCOUNTNUMBER")
    c_chq = pick("CHEQUENO", "CHEQUENUMBER")
    c_cdt = pick("CHEQUEDATE")
    c_amt = pick("CLAIMAMOUNT", "CHEQUEAMOUNT")
    c_bank = pick("BANK", "DRAWNON")
    c_rsn = pick("CHEQUEREALISTAION", "CHEQUEREALISATION", "RETURNREASON")
    c_bn = pick("BORROWER", "BORROWERNAME")
    c_ba = pick("BORROWERADD", "BORROWERADDRESS")
    if not (c_agr and c_bn and c_ba):
        raise ValueError(f"Agreement / Borrower columns not found. Columns: {list(df.columns)}")

    co = {}
    for sq, orig in cols.items():
        m = re.fullmatch(r"COBORROWER(\d+)(NAME)?", sq)
        if m:
            co.setdefault(int(m.group(1)), {})["name"] = orig
        m = re.fullmatch(r"COBORROWER(\d+)ADD(?:RESS)?", sq)
        if m:
            co.setdefault(int(m.group(1)), {})["add"] = orig
    co = {n: v for n, v in sorted(co.items()) if "name" in v}

    agreements = {}
    for _, r in df.iterrows():
        agr = "" if isnull(r[c_agr]) else str(r[c_agr]).strip().upper()
        if not agr:
            continue
        if agr in agreements:
            print(f"  WARN: agreement {agr} appears more than once in company data - using first row")
            continue
        parties = []
        if not isnull(r[c_bn]) and str(r[c_bn]).strip():
            parties.append({"slot": "Borrower", "name": r[c_bn], "addr": r[c_ba]})
        for n, v in co.items():
            if isnull(r[v["name"]]) or not str(r[v["name"]]).strip():
                continue
            addr = r[v["add"]] if "add" in v else ""
            parties.append({"slot": f"Co-Borrower-{n}", "name": r[v["name"]], "addr": addr})
        g = lambda c: (r[c] if c else None)
        agreements[agr] = {
            "parties": parties,
            "cheque_no": g(c_chq),
            "cheque_date": g(c_cdt),
            "amount": g(c_amt),
            "bank": g(c_bank),
            "reason": g(c_rsn),
        }
    return agreements


# ----------------------------------------------------------------------------
# PART 3 - RECONCILIATION
# ----------------------------------------------------------------------------
def party_status(name_sim, addr_sim):
    if name_sim >= MATCH_T and addr_sim >= MATCH_T:
        return "MATCH"
    if name_sim >= PARTIAL_T and addr_sim >= PARTIAL_T:
        return "PARTIAL MATCH"
    return "MISMATCH"


def match_parties(notice_rows, company_parties):
    """
    Order-independent matching. Every (notice entry, company entry) pair is
    scored on name (60%) + address (40%); best pairs are locked in first, so
    - the same person listed twice (two addresses) pairs address-to-address,
    - a party missing on either side does not shift everyone else out of line.
    """
    pairs = []
    for i, n in enumerate(notice_rows):
        for j, c in enumerate(company_parties):
            ns = similarity(n["Name"], c["name"])
            a_s = similarity(n["Address"], c["addr"])
            if ns >= PAIR_MIN_NAME:
                pairs.append((0.6 * ns + 0.4 * a_s, i, j, ns, a_s))
    pairs.sort(reverse=True)
    used_i, used_j, assign = set(), set(), {}
    for _, i, j, ns, a_s in pairs:
        if i in used_i or j in used_j:
            continue
        used_i.add(i)
        used_j.add(j)
        assign[i] = (j, ns, a_s)
    return assign


def cheque_checks(first, comp):
    """Returns list of (field, notice_value, company_value, result)."""
    out = []

    def add(field, nv, cv, ok):
        out.append((field, nv, cv, ok))

    # cheque no
    nv = first["Cheque No"]
    cv = comp["cheque_no"]
    if not nv:
        res = "NOT EXTRACTED"
    else:
        res = "OK" if cheque_key(nv) == cheque_key(cv) else "MISMATCH"
    add("Cheque No", nv, "" if isnull(cv) else cv, res)

    # cheque date
    nd, cd = parse_date(first["Cheque Date"]), parse_date(comp["cheque_date"])
    res = "NOT EXTRACTED" if nd is None else ("OK" if nd == cd else "MISMATCH")
    add("Cheque Date", nd, cd, res)

    # amount
    na, ca = to_number(first["Cheque Amount"]), to_number(comp["amount"])
    res = "NOT EXTRACTED" if na is None else ("OK" if ca is not None and abs(na - ca) < 0.5 else "MISMATCH")
    add("Cheque Amount", na, ca, res)

    # bank
    nb, cb = norm(first["Drawn On"]), norm(comp["bank"])
    if not nb:
        res = "NOT EXTRACTED"
    else:
        res = "OK" if (nb in cb or cb in nb) and cb else ("OK" if similarity(nb, cb) >= 85 else "MISMATCH")
    add("Drawn On / Bank", first["Drawn On"], "" if isnull(comp["bank"]) else comp["bank"], res)

    # return reason (compare code such as F009 if present, else the text)
    nr, cr = first["Return Reason"], comp["reason"]
    if not nr:
        res = "NOT EXTRACTED"
    else:
        code = lambda s: (re.match(r"\s*([A-Z]\d{3})\b", str(s).upper()) or [None, None])[1]
        nc, cc = code(nr), code(cr)
        if nc and cc:
            res = "OK" if nc == cc else "MISMATCH"
        else:
            res = "OK" if similarity(nr, cr) >= 85 else "MISMATCH"
    add("Return Reason", nr, "" if isnull(cr) else cr, res)
    return out


def reconcile(extracted, companies):
    party_rows, summary_rows = [], []

    for (src, agr), grp in extracted.groupby(["Source File", "Agreement No"], sort=False):
        rows = grp.to_dict("records")
        comp = companies.get(agr)

        if comp is None:
            for r in rows:
                party_rows.append(_party_row(r, None, "NOT IN COMPANY DATA", None, None,
                                             "Agreement number not found in company file."))
            summary_rows.append({"Source File": src, "Agreement No": agr, "Verdict": "NOT FOUND",
                                 "Issues": "Agreement not found in company data"})
            continue

        cparties = comp["parties"]
        assign = match_parties(rows, cparties)
        matched_j = {j for j, _, _ in assign.values()}
        issues = []

        for i, r in enumerate(rows):
            if i not in assign:
                party_rows.append(_party_row(r, None, "NOT IN COMPANY DATA", None, None,
                                             "No company borrower/co-borrower has a similar name."))
                issues.append(f"Party {r['Party Number']} ({r['Name']}) not in company data")
                continue
            j, ns, a_s = assign[i]
            c = cparties[j]
            st = party_status(ns, a_s)
            party_rows.append(_party_row(r, c, st, ns, a_s, _remark(st, ns, a_s)))
            if st != "MATCH":
                issues.append(f"Party {r['Party Number']} addr {r['Address Seq']} ({r['Name']}): {st}")

        # company parties that never appeared in the notice
        notice_names = [(norm(rows[i]["Name"]), i) for i in assign]
        for j, c in enumerate(cparties):
            if j in matched_j:
                continue
            same = [i for (nn, i) in notice_names if similarity(nn, c["name"]) >= MATCH_T]
            if same:
                rm = (f"Company lists a further address for '{rows[same[0]]['Name']}' that is not in "
                      "the notice (add it as 'Also at ...' or confirm it was dropped).")
            else:
                rm = "Party is in company data but missing from the notice."
            party_rows.append({
                "Source File": src, "Agreement No": agr, "Notice Party No": "", "Address Seq": "",
                "Status": "MISSING IN NOTICE", "Notice Name": "", "Company Name": c["name"],
                "Name Similarity %": None, "Notice Address": "", "Company Address": c["addr"],
                "Address Similarity %": None, "Company Slot": c["slot"], "Remarks": rm,
            })
            issues.append(f"{c['slot']} ({c['name']}) missing from notice")

        checks = cheque_checks(rows[0], comp)
        srow = {"Source File": src, "Agreement No": agr}
        for field, nv, cv, res in checks:
            srow[f"{field} - Notice"] = nv
            srow[f"{field} - Company"] = cv
            srow[f"{field} - Result"] = res
            if res != "OK":
                issues.append(f"{field}: {res}")
        srow["Notice Entries"] = len(rows)
        srow["Company Parties"] = len(cparties)
        srow["Verdict"] = "OK" if not issues else "CHECK"
        srow["Issues"] = "; ".join(issues)
        summary_rows.append(srow)

    p = pd.DataFrame(party_rows)
    s = pd.DataFrame(summary_rows)
    lead = ["Source File", "Agreement No", "Verdict", "Issues"]
    s = s[lead + [c for c in s.columns if c not in lead]]
    return p, s


def _party_row(r, c, status, ns, a_s, remark):
    return {
        "Source File": r["Source File"], "Agreement No": r["Agreement No"],
        "Notice Party No": r["Party Number"], "Address Seq": r["Address Seq"],
        "Status": status, "Notice Name": r["Name"],
        "Company Name": c["name"] if c else "",
        "Name Similarity %": None if ns is None else round(ns, 1),
        "Notice Address": r["Address"], "Company Address": c["addr"] if c else "",
        "Address Similarity %": None if a_s is None else round(a_s, 1),
        "Company Slot": c["slot"] if c else "", "Remarks": remark,
    }


def _remark(status, ns, a_s):
    if status == "MATCH":
        return "Name and address match."
    bits = []
    if ns < MATCH_T:
        bits.append(f"name differs ({ns:.0f}%)")
    if a_s < MATCH_T:
        bits.append(f"address differs ({a_s:.0f}%)")
    return "Check: " + " and ".join(bits) + "."


# ----------------------------------------------------------------------------
# PART 4 - OUTPUT
# ----------------------------------------------------------------------------
def style_and_save(sheets, path):
    from openpyxl.styles import Alignment, Font, PatternFill

    green = PatternFill("solid", fgColor="C6EFCE")
    amber = PatternFill("solid", fgColor="FFEB9C")
    red = PatternFill("solid", fgColor="FFC7CE")
    colour = {"MATCH": green, "OK": green, "PARTIAL MATCH": amber, "CHECK": amber,
              "NOT EXTRACTED": amber, "MISSING IN NOTICE": amber, "MISMATCH": red,
              "NOT IN COMPANY DATA": red, "NOT FOUND": red}

    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        for name, df in sheets.items():
            df.to_excel(xw, sheet_name=name, index=False)
            ws = xw.sheets[name]
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            for i, col in enumerate(df.columns, start=1):
                width = max([len(str(col))] + [len(str(v)) for v in df[col].head(300)])
                ws.column_dimensions[ws.cell(1, i).column_letter].width = min(max(12, width + 2), 55)
                ws.cell(1, i).font = Font(bold=True)
            for row in ws.iter_rows(min_row=2):
                for cell in row:
                    cell.alignment = Alignment(vertical="top", wrap_text=isinstance(cell.value, str) and len(cell.value) > 40)
                    if isinstance(cell.value, str) and cell.value in colour:
                        cell.fill = colour[cell.value]


def save_safely(fn, path):
    try:
        fn(path)
        return path
    except PermissionError:
        alt = path.replace(".xlsx", f"_{int(time.time())}.xlsx")
        fn(alt)
        print(f"NOTE: '{path}' is open in Excel - saved as {alt}")
        return alt


# ----------------------------------------------------------------------------
# MAIN
# ----------------------------------------------------------------------------
def main():
    here = os.path.dirname(os.path.abspath(__file__))
    ap = argparse.ArgumentParser(description="Extract notice PDFs and reconcile with company data")
    ap.add_argument("--pdf-folder", default=here)
    ap.add_argument("--company", default=os.path.join(here, "company.xlsx"))
    ap.add_argument("--out", default=os.path.join(here, "reconciliation_report.xlsx"))
    ap.add_argument("--extracted-out", default=os.path.join(here, "extracted_notices.xlsx"))
    args = ap.parse_args()

    pdfs = sorted(f for f in os.listdir(args.pdf_folder) if f.lower().endswith(".pdf"))
    print(f"Found {len(pdfs)} PDF file(s)\n")

    data = []
    for f in pdfs:
        try:
            recs = extract_data_from_pdf(os.path.join(args.pdf_folder, f))
            if not recs:
                print(f"WARN  [{f}]: no recipients found - check the 'To ...' block")
                continue
            data.extend(recs)
            parties = len({r["Party Number"] for r in recs})
            extra = len(recs) - parties
            print(f"OK    [{f}]: {parties} parties, {len(recs)} address rows"
                  + (f" ({extra} 'Also at' address(es))" if extra else ""))
        except Exception as e:
            print(f"ERROR [{f}]: {e}")

    if not data:
        sys.exit("\nNo data extracted.")

    extracted = pd.DataFrame(data)
    save_safely(lambda p: extracted.to_excel(p, index=False), args.extracted_out)

    print(f"\nLoading company data: {args.company}")
    companies = load_company(args.company)
    print(f"  {len(companies)} agreements loaded")

    parties, summary = reconcile(extracted, companies)
    out = save_safely(
        lambda p: style_and_save(
            {"Summary": summary, "Party Comparison": parties, "Extracted": extracted}, p),
        args.out,
    )

    print("\nVerdicts:")
    print(summary["Verdict"].value_counts().to_string())
    print(f"\nReport: {out}")


if __name__ == "__main__":
    main()
